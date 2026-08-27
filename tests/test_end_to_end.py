"""A full pipeline run with the network and the AI stubbed.

This is the test that would have caught the missing-Flask bug in ExhibitorPilot
had it existed there: it exercises every stage's wiring, not just its helpers.
"""

import json
import re

import pytest


SITE_HTML = """<html><head><title>Acme Mining Equipment</title></head><body>
<a href="/global-presence">Global Presence</a><a href="/about">About</a>
<main><h1>Acme</h1>
<p>We manufacture jaw crushers and screens for mining and quarrying.</p>
<p>Our Australian subsidiary Acme Pty Ltd is based in Perth, Western Australia.</p>
<p>Established 1985. 450 employees. Plant in Pune, India.</p>
<p>Managing Director: Ravi Sharma. Export Director: Priya Nair.</p>
</main></body></html>"""


class FakeResult:
    def __init__(self, title, url, snippet):
        self.title, self.url, self.snippet = title, url, snippet


class FakeLLM:
    """Answers shaped by which stage's prompt it was handed."""

    def __init__(self):
        self.calls = 0

    def ask_json(self, prompt, system="", max_tokens=0, temperature=0.0):
        self.calls += 1
        if "Extract company names" in prompt:
            return {"companies": [
                {"name": "Acme Crushers Ltd", "source_url": "https://list.example/top"},
                {"name": "Bharat Screening Works", "source_url": "https://list.example/top"},
                {"name": "10 Best Crusher Makers", "source_url": ""},   # must be filtered
            ]}
        if "Classify this company" in prompt:
            return {"category": "Crushing and screening",
                    "products": "Jaw crushers and vibrating screens",
                    "entity_type": "Manufacturer", "country": "India",
                    "origin_country": "India", "relevance": "High",
                    "reasoning": "Core crushing equipment.", "from_name_only": False}
        if "Judge this company" in prompt:
            return {"level": "Strong match",
                    "matched_criteria": ["Sells into Australia"],
                    "evidence": ["Australian subsidiary - Acme Pty Ltd, Perth WA"],
                    "sources": ["https://acme.example/global-presence"],
                    "summary": "Owns an Australian subsidiary."}
        if "Extract basic company information" in prompt:
            return {"revenue": "Not publicly available", "employees": "450",
                    "founded": "1985", "ownership": "Private", "parent": "Independent",
                    "plants": "Pune, India", "export_countries": "Australia"}
        if "Write the first contact message" in prompt:
            return {"angle": "Australian subsidiary in Perth",
                    "subject": "parts lead times into perth",
                    "body": "I saw Acme runs its own Perth arm. Are parts lead "
                            "times from India a constraint for your customers there?",
                    "usable": True}
        if "Identify the people" in prompt:
            return {"people": [
                {"name": "Ravi Sharma", "title": "Managing Director", "linkedin": "",
                 "note": "Decision maker"},
                {"name": "Priya Nair", "title": "Export Director",
                 "linkedin": "https://linkedin.com/in/example", "note": "Owns exports"},
                {"name": "Made Up", "title": "Intern",
                 "linkedin": "http://not-linkedin.example/x", "note": ""},
            ]}
        raise AssertionError(f"unexpected prompt: {prompt[:120]}")

    def chat(self, *a, **k):
        return "OK"


@pytest.fixture
def stubbed(monkeypatch, sample_plan):
    from prospector import config, fetcher, planner
    from prospector.stages import (
        classify, discover, outreach, people, profile, qualify, resolve)

    # The Drafter refuses without a sender profile, by design.
    sample_plan["sender_profile"] = ("I place Indian crushing equipment with "
                                     "Australian quarry groups.")
    config.save_plan(sample_plan)
    fake = FakeLLM()

    # `planner` is patched too, and deliberately: discover falls back to
    # expand_queries when it runs short of the target, and that call would
    # otherwise reach the real API over the network from inside a test.
    for module in (classify, qualify, profile, people, discover, outreach, planner):
        monkeypatch.setattr(module, "get_client", lambda stage="default": fake)

    monkeypatch.setattr(fetcher, "fetch_html", lambda url, use_cache=True: SITE_HTML)

    def fake_search(query, max_results=8, use_cache=True):
        # Derive the domain from the query so resolve's real scoring logic is
        # genuinely exercised rather than bypassed by a constant URL.
        name = query.split('"')[1] if '"' in query else query
        tokens = [t for t in re.findall(r"[a-z0-9]+", name.lower()) if len(t) > 2]
        slug = tokens[0] if tokens else "acme"
        return [FakeResult(f"{name} - Official Website", f"https://{slug}.example/",
                           "Equipment for mining.")]

    monkeypatch.setattr(resolve, "search", fake_search)
    monkeypatch.setattr(discover, "search", fake_search)
    for module in (qualify, profile, people):
        monkeypatch.setattr(module, "search_many",
                            lambda queries, max_results=6: [
                                FakeResult("Acme Pty Ltd Australia",
                                           "https://acme.example/au",
                                           "Perth WA subsidiary.")])
    return fake


def test_full_run_from_prompt_to_spreadsheet(stubbed, tmp_path):
    from prospector.database import get_stats
    from prospector.pipeline import run_pipeline

    out = tmp_path / "leads.xlsx"
    result = run_pipeline(workers=2, export_path=out, quiet=True)

    assert result["errors"] == {}, result["errors"]

    stats = get_stats()
    assert stats["total"] >= 2          # discovered, with the article title filtered out
    assert stats["resolved"] == stats["total"]
    assert stats["classified"] == stats["total"]
    assert stats["qualified"] == stats["total"]
    assert stats["with_people"] == stats["total"]
    assert stats["drafted"] == stats["total"]
    assert out.exists()

    from openpyxl import load_workbook
    wb = load_workbook(out)
    assert wb.sheetnames == ["Read Me", "Call List", "Qualification",
                             "Company Profiles", "Contacts", "First Contact",
                             "All Leads", "Summary"]
    # The draft has to survive into the workbook, or the stage was pointless.
    assert "Perth" in str(wb["First Contact"].cell(row=4, column=6).value)
    assert wb["Call List"].max_row > 3


def test_article_titles_never_reach_the_database(stubbed):
    from prospector.database import get_connection
    from prospector.pipeline import run_pipeline

    run_pipeline(stages=["discover"], quiet=True)
    names = [r[0] for r in get_connection().execute(
        "SELECT company_name FROM leads").fetchall()]
    assert "Acme Crushers Ltd" in names
    assert not any(n.startswith("10 Best") for n in names)


def test_the_run_is_resumable(stubbed):
    from prospector.database import get_stats
    from prospector.pipeline import run_pipeline

    run_pipeline(stages=["discover", "resolve"], quiet=True)
    first = get_stats()["resolved"]
    assert first > 0

    # Running again must not redo finished rows.
    run_pipeline(stages=["resolve"], quiet=True)
    assert get_stats()["resolved"] == first


def test_invented_linkedin_urls_are_stripped_and_roles_ranked(stubbed):
    from prospector.database import get_connection
    from prospector.pipeline import run_pipeline

    run_pipeline(stages=["discover", "resolve", "crawl", "classify", "qualify", "people"],
                 quiet=True)
    row = get_connection().execute(
        "SELECT people_json FROM leads WHERE people_json IS NOT NULL").fetchone()
    people = json.loads(row["people_json"])

    assert len(people) <= 3
    assert all("linkedin.com" in p["linkedin"] or p["linkedin"] == "" for p in people)
    assert people[0]["title"] == "Managing Director"


def test_qualification_score_and_level_are_stored(stubbed):
    from prospector.database import get_connection
    from prospector.pipeline import run_pipeline

    run_pipeline(stages=["discover", "resolve", "crawl", "classify", "qualify"], quiet=True)
    row = get_connection().execute(
        "SELECT qualification_level, qualification_score, qualification_matched "
        "FROM leads WHERE qualification_level IS NOT NULL").fetchone()

    assert row["qualification_level"] == "Strong match"
    assert row["qualification_score"] > 70
    assert "Australia" in row["qualification_matched"]


def test_preflight_blocks_a_run_with_no_plan():
    from prospector.pipeline import preflight

    problems = preflight(["discover"])
    assert problems and "No research plan" in problems[0]


def test_preflight_blocks_a_run_with_no_ai_at_all(monkeypatch, sample_plan):
    from prospector import config
    from prospector.pipeline import preflight

    config.save_plan(sample_plan)
    monkeypatch.delenv("OPENROUTER_API_KEY", raising=False)

    monkeypatch.delenv("LOCAL_MODEL", raising=False)

    problems = preflight(["classify"])
    assert problems and "No AI set up" in problems[0]


def test_the_spreadsheet_says_what_each_row_was_researched_from(stubbed, tmp_path):
    """A name-only guess and a fully-read website used to look identical.

    The Judge is handed the guessed products string as though it were fact, so
    the reader needs to know which kind of row they are looking at.
    """
    from openpyxl import load_workbook
    from prospector.pipeline import run_pipeline

    out = tmp_path / "leads.xlsx"
    run_pipeline(workers=2, export_path=out, quiet=True)

    ws = load_workbook(out)["All Leads"]
    headers = [c.value for c in ws[1]]
    assert "Researched from" in headers

    column = headers.index("Researched from") + 1
    assert ws.cell(row=2, column=column).value


def test_the_summary_counts_reconcile_with_the_lead_count(stubbed, tmp_path):
    """Unclassified rows vanished from the breakdown with no residual line."""
    from openpyxl import load_workbook
    from prospector.pipeline import run_pipeline

    out = tmp_path / "leads.xlsx"
    run_pipeline(workers=2, export_path=out, quiet=True)

    ws = load_workbook(out)["Summary"]
    labels = [ws.cell(row=r, column=1).value for r in range(1, ws.max_row + 1)]
    assert any("Not yet classified" in str(l) for l in labels)


def test_a_blocked_search_does_not_fabricate_qualification_ratings(monkeypatch, sample_plan):
    """When search is blocked and no site was read, there is nothing to judge.

    Left alone the model was asked to rate a company from its name and the
    prompt boilerplate, and every remaining row filled with confident-looking
    "Partial match" ratings built on nothing.
    """
    from prospector import config
    from prospector.stages import qualify

    config.save_plan(sample_plan)

    called = []
    monkeypatch.setattr(qualify, "get_client",
                        lambda stage="default": called.append(1))
    monkeypatch.setattr(qualify, "_site_evidence",
                        lambda row, kw, max_chars=6000: ("(no website content available)", []))
    monkeypatch.setattr(qualify, "_search_evidence",
                        lambda name, plan, max_chars=6000: ("(no search results)", []))

    row = {"company_name": "Acme Ltd", "website": "", "products": "", "pages_json": None}

    # Raised, not returned as a verdict. Writing "Unclear" here would stamp
    # `qualified_at` and take the row out of the pending query for good, so a
    # ten-minute search outage would become a permanent answer -- and pressing
    # Continue, which is exactly what the app tells the user to do, would skip
    # every company it had touched.
    with pytest.raises(ValueError, match="tried again"):
        qualify.qualify_one(row, sample_plan)

    assert not called, "the model must not be asked to judge with no sources"
